import json
import logging
import os
import operator
import re
from typing import Annotated, Any, TypedDict

import httpx
from dotenv import load_dotenv
from pydantic import BaseModel, Field

from langchain_core.messages import AIMessage, HumanMessage, SystemMessage, ToolMessage
from langchain_core.rate_limiters import InMemoryRateLimiter
from langchain_mistralai import ChatMistralAI
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.graph import END, START, StateGraph
from langgraph.graph.message import add_messages
from langgraph.prebuilt import ToolNode
from langgraph.types import Command, RetryPolicy, interrupt

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from tools import search_semantic_scholar

load_dotenv()

# Silence the harmless Gemini AFC warning
logging.getLogger("google_genai.models").setLevel(logging.ERROR)
import warnings
warnings.filterwarnings("ignore", message=".*fixed sampling defaults.*")


# CONFIG


MAX_REVISIONS = int(os.getenv("MAX_REVISIONS", "2"))
MAX_TOOL_MESSAGES = 8          # cap on Semantic Scholar tool rounds
MIN_RELEVANT_SOURCES = 15
MAX_CHARS = int(os.getenv("MAX_PROMPT_CHARS", "14000"))
MANUSCRIPT_CHARS = 40000


# MODELS  (rate limited + retries + optional Gemini fallback)


from langchain_core.callbacks import BaseCallbackHandler


class LLMErrorLogger(BaseCallbackHandler):
    """Prints the real error from each provider (with_fallbacks hides all but the first)."""

    def __init__(self, name: str):
        self.name = name

    def on_llm_error(self, error, **kwargs):
        print(f"[{self.name} error] {type(error).__name__}: {str(error)[:200]}")


PRIMARY = os.getenv("PRIMARY_LLM", "mistral").lower()

mistral = ChatMistralAI(
    model="mistral-small-latest",
    temperature=0.2,
    rate_limiter=InMemoryRateLimiter(
        requests_per_second=float(os.getenv("MISTRAL_RPS", "0.2")),
        check_every_n_seconds=0.1,
        max_bucket_size=1,
    ),
    max_retries=2,
    callbacks=[LLMErrorLogger("mistral")],
)

gemini = None
if os.getenv("GOOGLE_API_KEY"):
    try:
        from langchain_google_genai import ChatGoogleGenerativeAI

        gemini = ChatGoogleGenerativeAI(
            model=os.getenv("GEMINI_MODEL", "gemini-3.5-flash-lite"),
            rate_limiter=InMemoryRateLimiter(
                requests_per_second=float(os.getenv("GEMINI_RPS", "0.2")),
                check_every_n_seconds=0.1,
                max_bucket_size=1,
            ),
            callbacks=[LLMErrorLogger("gemini")],
        )
    except Exception as exc:  # noqa: BLE001
        print(f"Gemini unavailable: {exc}")

# Tool-calling model for literature search
tools = [search_semantic_scholar]
llm_with_tools = (gemini if gemini is not None else mistral).bind_tools(tools)


def structured(schema):
    """Structured-output chain: PRIMARY_LLM first, the other provider as fallback."""
    mistral_chain = mistral.with_structured_output(schema)
    if gemini is None:
        return mistral_chain
    gemini_chain = gemini.with_structured_output(schema)
    if PRIMARY == "gemini":
        return gemini_chain.with_fallbacks([mistral_chain])
    return mistral_chain.with_fallbacks([gemini_chain])


def _should_retry(exc: Exception) -> bool:
    if isinstance(exc, httpx.HTTPStatusError):
        return exc.response.status_code in (429, 500, 502, 503, 504)
    return isinstance(exc, (httpx.TransportError, ConnectionError, TimeoutError))


RETRY = RetryPolicy(
    max_attempts=3,
    initial_interval=5.0,
    backoff_factor=2.0,
    max_interval=60.0,
    retry_on=_should_retry,
)


# HELPERS



def _txt(x: Any, limit: int = MAX_CHARS) -> str:
    """Stringify and truncate so prompts stay within token limits."""
    s = x if isinstance(x, str) else json.dumps(x, indent=1, ensure_ascii=False, default=str)
    return s if len(s) <= limit else s[:limit] + "\n...[truncated]"


def run_agent(llm, system: str, parts: dict[str, Any], limit: int = MAX_CHARS) -> dict:
    human = "\n\n".join(f"===== {k} =====\n{_txt(v, limit)}" for k, v in parts.items())
    result = llm.invoke([SystemMessage(content=system), HumanMessage(content=human)])
    return result.model_dump()


# NOTE: "Experimental Setup" has been removed as a standalone section, and
# "Results" + "Discussion" have been merged into a single "Result and
# Discussion" section (key: result_and_discussion).
SECTION_ORDER = [
    ("Introduction", "introduction"),
    ("Literature Review", "literature_review"),
    ("Methodology", "methodology"),
    ("Result and Discussion", "result_and_discussion"),
    ("Conclusion", "conclusion"),
    ("Future Scope", "future_scope"),
]
SECTION_KEYS = [k for _, k in SECTION_ORDER]

_KEY_MAP = {
    "abstract": "abstract",
    "introduction": "introduction",
    "literature review": "literature_review",
    "literature_review": "literature_review",
    "related work": "literature_review",
    "related_work": "literature_review",
    "background": "literature_review",
    "methodology": "methodology",
    "methods": "methodology",
    "method": "methodology",
    "materials and methods": "methodology",
    "results": "result_and_discussion",
    "result": "result_and_discussion",
    "results and analysis": "result_and_discussion",
    "discussion": "result_and_discussion",
    "results and discussion": "result_and_discussion",
    "result and discussion": "result_and_discussion",
    "results_discussion": "result_and_discussion",
    "result_and_discussion": "result_and_discussion",
    "conclusion": "conclusion",
    "conclusions": "conclusion",
    "future scope": "future_scope",
    "future work": "future_scope",
    "future_scope": "future_scope",
}


def section_key(title: str) -> str:
    """Map a section title ('3. Methodology', 'III. Methods') to a state key ('' if none)."""
    t = re.sub(r"^\s*([0-9]+|[ivxIVX]+)[\.\)]\s*", "", title or "").strip().lower()
    return _KEY_MAP.get(t, "")


def build_manuscript(state: dict, limit: int = MANUSCRIPT_CHARS) -> str:
    parts = []
    for title, key in [("Abstract", "abstract")] + SECTION_ORDER:
        content = state.get(key, "")
        if content:
            parts.append(f"===== {title} =====\n\n{content}\n")
    if not parts:
        return "No manuscript sections are currently available."
    return _txt("\n".join(parts), limit)


def _read_uploaded(files: list[str], per_file: int = 8000) -> str:
    out = []
    for path in files or []:
        if os.path.isfile(path) and path.lower().endswith(
            (".txt", ".md", ".csv", ".json", ".py", ".tex")
        ):
            try:
                with open(path, "r", encoding="utf-8", errors="ignore") as f:
                    out.append(f"--- {path} ---\n{f.read()[:per_file]}")
                continue
            except OSError:
                pass
        out.append(f"--- {path} (content not read) ---")
    return "\n\n".join(out)


def literature_brief(state: dict, abstract_chars: int = 500) -> list[dict]:
    brief = []
    for p in state.get("literature_results", []):
        brief.append(
            {
                "source_id": p.get("source_id"),
                "title": p.get("title"),
                "authors": p.get("authors"),
                "year": p.get("year"),
                "doi": p.get("doi"),
                "venue": p.get("venue"),
                "volume": p.get("volume"),
                "issue": p.get("issue"),
                "pages": p.get("pages"),
                "url": p.get("url"),
                "citation_count": p.get("citation_count", 0),
                "abstract": (p.get("abstract") or "")[:abstract_chars],
            }
        )
    return brief


NO_INVENTION = """
GLOBAL RULES (apply to everything below):
- Use ONLY information present in the supplied input. Never invent facts, numbers,
  datasets, algorithms, metrics, results, papers, authors, DOIs, URLs or citations.
- If information is missing, say it is missing; do not fill it with assumptions.
- Do not exaggerate novelty.
"""


# STATE



class ResearchState(TypedDict, total=False):
    # Input
    project_description: str
    uploaded_files: list[str]
    messages: Annotated[list, add_messages]

    # Understanding / planning
    project_profile: dict[str, Any]
    missing_information: list[str]
    confirmed_facts: list[str]
    uncertain_facts: list[str]
    research_problem: str
    research_objectives: list[str]
    research_contributions: list[str]
    research_plan: dict[str, Any]
    paper_outline: dict[str, Any]

    # Literature
    literature_results: list[dict[str, Any]]
    literature_matrix: list[dict[str, Any]]
    literature_synthesis: dict[str, Any]
    research_gap: dict[str, Any]

    # Human approval
    approval_status: str
    approval_feedback: str

    # Writing
    completed_sections: list[str]
    title_candidates: list[str]
    selected_title: str
    abstract: str
    index_terms: list[str]
    introduction: str
    literature_review: str
    methodology: str
    result_and_discussion: str
    conclusion: str
    future_scope: str
    references: list[dict[str, Any]]

    # Verified experimental data supplied by the user
    actual_results: dict[str, Any]
    paper_metadata: dict[str, Any]

    # Verification
    citation_feedback: str
    claim_feedback: str
    consistency_feedback: str
    # reducer: verification history accumulates instead of being overwritten
    verification_results: Annotated[list[dict[str, Any]], operator.add]
    citation_verification_results: list[dict[str, Any]]
    citation_records: list[dict[str, Any]]
    citation_issues: list[dict[str, Any]]
    claim_issues: list[dict[str, Any]]
    consistency_issues: list[dict[str, Any]]
    review_issues: list[dict[str, Any]]

    # Review / control
    review_feedback: str
    revision_count: int
    is_approved: bool
    current_stage: str
    final_validation_status: str
    final_validation_issues: list[str]
    final_validation_warnings: list[str]

    # Output
    final_manuscript: str
    output_files: list[str]



# 1. SUPERVISOR



def supervisor_node(state: ResearchState) -> dict:
    print("\n--- [Stage 1] Supervisor ---")
    return {"current_stage": "onboarding", "revision_count": state.get("revision_count", 0)}



# 2. PROJECT UNDERSTANDING



class ProjectProfile(BaseModel):
    research_title: str = ""
    research_problem: str = ""
    domain: str = ""
    motivation: str = ""
    objective: str = ""
    proposed_solution: str = ""
    methodology: str = ""
    dataset: str = ""
    dataset_size: str = ""
    input_output: str = ""
    algorithms_models: list[str] = Field(default_factory=list)
    frameworks: list[str] = Field(default_factory=list)
    preprocessing: str = ""
    training_methodology: str = ""
    experimental_setup: str = ""
    evaluation_metrics: list[str] = Field(default_factory=list)
    actual_results: str = ""
    expected_contribution: str = ""
    limitations: str = ""
    hardware: str = ""
    software: str = ""
    deployment_details: str = ""
    confirmed_facts: list[str] = Field(default_factory=list)
    uncertain_facts: list[str] = Field(default_factory=list)
    missing_information: list[str] = Field(default_factory=list)


project_understanding_llm = structured(ProjectProfile)

PROJECT_UNDERSTANDING_PROMPT = NO_INVENTION + """
You are the Project Understanding Agent of PaperForge AI.
Understand the user's research project before any planning or writing begins.
Extract only explicitly available information about: title/topic, research problem,
domain, motivation, objective, proposed solution, methodology, dataset and size,
input/output, algorithms/models, frameworks, preprocessing, training methodology,
experimental setup (datasets, splits, hardware, software, hyperparameters,
baselines, evaluation protocol - to be folded into Result and Discussion later),
evaluation metrics, actual results, expected contribution, limitations, hardware,
software, and deployment details.
Separate information into:
- confirmed_facts: clearly provided by the user
- uncertain_facts: unclear or ambiguous
- missing_information: required for a paper but not provided
Leave unknown fields empty.
"""


def project_understanding_node(state: ResearchState) -> dict:
    print("\n--- [Stage 2] Project Understanding ---")
    messages = state.get("messages", [])
    latest = messages[-1].content if messages and hasattr(messages[-1], "content") else ""
    profile = run_agent(
        project_understanding_llm,
        PROJECT_UNDERSTANDING_PROMPT,
        {
            "USER PROJECT DESCRIPTION": state.get("project_description", ""),
            "LATEST USER MESSAGE": latest,
            "UPLOADED PROJECT FILES": _read_uploaded(state.get("uploaded_files", [])),
        },
    )
    return {
        "project_profile": profile,
        "research_problem": profile["research_problem"],
        "research_objectives": [profile["objective"]] if profile["objective"] else [],
        "missing_information": profile["missing_information"],
        "confirmed_facts": profile["confirmed_facts"],
        "uncertain_facts": profile["uncertain_facts"],
    }



# 3. RESEARCH PLANNER



class ResearchPlan(BaseModel):
    research_problem: str = ""
    research_objectives: list[str] = Field(default_factory=list)
    research_questions: list[str] = Field(default_factory=list)
    research_contributions: list[str] = Field(default_factory=list)
    methodology_approach: str = ""
    evaluation_strategy: str = ""
    literature_topics: list[str] = Field(default_factory=list)
    required_evidence: list[str] = Field(default_factory=list)
    limitations: list[str] = Field(default_factory=list)


research_planner_llm = structured(ResearchPlan)

RESEARCH_PLANNER_PROMPT = NO_INVENTION + """
You are the Research Planner Agent of PaperForge AI.
Create a research plan based ONLY on the project information provided.
Determine: the research problem, objectives, research questions, expected
contributions, methodology (only what is already described), evaluation strategy,
literature topics to investigate, evidence required, and stated limitations.
- Do not claim novelty and do not identify a research gap (a later agent does that).
- Do not search for papers.
- Make literature_topics and research_questions specific enough to generate good
  academic search queries.
- If HUMAN FEEDBACK on a previous outline is provided, revise the plan to address it.
"""


def research_planner_node(state: ResearchState) -> dict:
    print("\n--- [Stage 3] Research Planner ---")
    plan = run_agent(
        research_planner_llm,
        RESEARCH_PLANNER_PROMPT,
        {
            "PROJECT PROFILE": state.get("project_profile", {}),
            "CONFIRMED FACTS": state.get("confirmed_facts", []),
            "UNCERTAIN FACTS": state.get("uncertain_facts", []),
            "MISSING INFORMATION": state.get("missing_information", []),
            "HUMAN FEEDBACK ON PREVIOUS OUTLINE": state.get("approval_feedback", ""),
        },
    )
    return {
        "research_problem": plan["research_problem"],
        "research_objectives": plan["research_objectives"],
        "research_contributions": plan["research_contributions"],
        "research_plan": plan,
        # fresh literature pass for the new plan
        "literature_results": [],
    }



# 4. LITERATURE RESEARCH (tool loop) + COLLECTOR


LITERATURE_RESEARCH_PROMPT = """
You are the Literature Research Agent of PaperForge AI.
Use the search_semantic_scholar tool to find academic papers relevant to the
research plan.
- Run 3 to 6 DIFFERENT, specific searches based on the literature topics and
  research questions (call the tool; do not answer from memory). Aim to collect
  at least 15 unique, relevant papers when the search results permit it.
- Review the tool results that come back. When you have enough relevant papers,
  stop calling tools and reply with a one-paragraph summary of what was found.
- Never invent papers, authors, DOIs, years or other metadata.
"""


def literature_research_node(state: ResearchState) -> dict:
    print("\n--- [Stage 4] Literature Research ---")
    prompt_input = (
        f"Research Problem:\n{state.get('research_problem', '')}\n\n"
        f"Research Objectives:\n{state.get('research_objectives', [])}\n\n"
        f"Research Plan:\n{_txt(state.get('research_plan', {}))}"
    )
    # Include accumulated messages so the model can see previous tool results.
    history = list(state.get("messages", []))
    messages = [SystemMessage(content=LITERATURE_RESEARCH_PROMPT), HumanMessage(content=prompt_input)] + history

    # Some providers (e.g. Gemini) reject a request whose final turn is an
    # assistant message with no tool call ("model prefilling" is not
    # supported; the final turn must be a user message or a function
    # response). route_after_literature can loop back to this node when the
    # model replied without calling the tool (to force more distinct
    # searches), which otherwise leaves `history` ending on a bare AIMessage.
    # Append a nudge HumanMessage in that case so every request ends on a
    # user turn, regardless of which provider is bound as llm_with_tools.
    if messages and isinstance(messages[-1], AIMessage) and not getattr(messages[-1], "tool_calls", None):
        messages.append(
            HumanMessage(
                content=(
                    "Run at least one more distinct search_semantic_scholar query "
                    "on a topic or angle you haven't tried yet before summarizing."
                )
            )
        )

    response = llm_with_tools.invoke(messages)
    return {"messages": [response]}


def route_after_literature(state: ResearchState) -> str:
    messages = state.get("messages", [])
    if not messages:
        return "collect_literature"
    tool_rounds = sum(1 for m in messages if isinstance(m, ToolMessage))
    agent_rounds = sum(1 for m in messages if not isinstance(m, ToolMessage))
    if getattr(messages[-1], "tool_calls", None) and tool_rounds < MAX_TOOL_MESSAGES:
        return "semantic_scholar"
    # Do not settle after one broad search. Re-prompt the retrieval agent for
    # distinct queries, while retaining a finite escape hatch for unavailable
    # tools or providers that decline to call the search tool.
    if tool_rounds < 3 and agent_rounds < 6:
        return "literature_research"
    return "collect_literature"


def _extract_papers(payload: Any) -> list[dict]:
    if isinstance(payload, list):
        items = payload
    elif isinstance(payload, dict):
        items = payload.get("data") or payload.get("papers") or payload.get("results") or [payload]
    else:
        return []
    return [p for p in items if isinstance(p, dict)]


def _normalise_paper(p: dict) -> dict:
    authors = p.get("authors", [])
    if isinstance(authors, list):
        authors = [a.get("name", "") if isinstance(a, dict) else str(a) for a in authors]
    elif authors:
        authors = [str(authors)]
    ext = p.get("externalIds") or {}
    publication_venue = p.get("publicationVenue") or {}
    journal = p.get("journal") or {}
    venue = p.get("venue") or ""
    if not venue and isinstance(publication_venue, dict):
        venue = publication_venue.get("name") or ""
    if not venue and isinstance(journal, dict):
        venue = journal.get("name") or ""
    volume = p.get("volume") or (journal.get("volume") if isinstance(journal, dict) else "") or ""
    issue = p.get("issue") or (journal.get("issue") if isinstance(journal, dict) else "") or ""
    pages = p.get("pages") or (journal.get("pages") if isinstance(journal, dict) else "") or ""
    return {
        **p,
        "title": p.get("title", ""),
        "authors": [a for a in authors if a],
        "year": str(p.get("year", "") or ""),
        "doi": p.get("doi") or (ext.get("DOI", "") if isinstance(ext, dict) else ""),
        "url": p.get("url", ""),
        "abstract": p.get("abstract") or "",
        "venue": venue,
        "volume": str(volume),
        "issue": str(issue),
        "pages": pages,
        "publication_date": p.get("publicationDate") or "",
        "citation_count": p.get("citationCount") or 0,
        "fields_of_study": p.get("fieldsOfStudy") or [],
    }


def _parse_tool_content(content: Any) -> Any:
    """Parse a tool result into JSON if possible (handles wrapped/embedded JSON)."""
    if isinstance(content, list):
        content = " ".join(c.get("text", "") if isinstance(c, dict) else str(c) for c in content)
    if not isinstance(content, str):
        return content
    try:
        return json.loads(content)
    except ValueError:
        pass
    # JSON embedded in surrounding text
    for opener, closer in (("[", "]"), ("{", "}")):
        i, j = content.find(opener), content.rfind(closer)
        if i != -1 and j > i:
            try:
                return json.loads(content[i : j + 1])
            except ValueError:
                continue
    return None


def collect_literature_node(state: ResearchState) -> dict:
    """Turn ToolMessages into structured literature_results + reference list."""
    print("\n--- [Stage 4b] Collect Literature ---")
    papers, seen = [], set()
    for n, m in enumerate(state.get("messages", [])):
        if not isinstance(m, ToolMessage):
            continue
        preview = m.content if isinstance(m.content, str) else str(m.content)
        print(f"[tool result {n}] status={getattr(m, 'status', '?')} :: {preview[:300]!r}")
        payload = _parse_tool_content(m.content)
        if payload is None:
            print("   -> could not parse this tool result as JSON")
            continue
        for raw in _extract_papers(payload):
            p = _normalise_paper(raw)
            key = (p.get("paperId") or p["doi"] or p["title"]).strip().lower()
            if not p["title"] or key in seen:
                continue
            seen.add(key)
            papers.append(p)

    for i, p in enumerate(papers, start=1):
        p["source_id"] = f"SRC{i}"

    references = [
        {k: p.get(k) for k in (
            "source_id", "title", "authors", "year", "doi", "url", "venue",
            "volume", "issue", "pages", "publication_date", "citation_count",
        )}
        for p in papers
    ]
    print(f"Collected {len(papers)} unique papers.")
    if 0 < len(papers) < MIN_RELEVANT_SOURCES:
        print(f"WARNING: only {len(papers)}/{MIN_RELEVANT_SOURCES} unique relevant papers were available.")
    if not papers:
        print("WARNING: no literature retrieved - citations will be unsupported. Check tools.py / Semantic Scholar limits.")
    return {"literature_results": papers, "references": references}



# 5. PAPER FILTERING, RANKING, ANALYSIS, AND SYNTHESIS



def paper_filtering_node(state: ResearchState) -> dict:
    """Keep uniquely identifiable, search-relevant sources and rank them."""
    print("\n--- [Stage 5] Paper Filtering and Ranking ---")
    query_terms = set(re.findall(r"[a-z]{4,}", " ".join([
        state.get("research_problem", ""),
        " ".join(state.get("research_objectives", [])),
        " ".join(state.get("research_plan", {}).get("literature_topics", [])),
    ]).lower()))
    ranked = []
    for paper in state.get("literature_results", []):
        complete = bool(
            paper.get("title") and paper.get("authors") and paper.get("year")
            and paper.get("abstract") and (paper.get("doi") or paper.get("url"))
        )
        if not complete:
            continue
        searchable = f"{paper.get('title', '')} {paper.get('abstract', '')}".lower()
        overlap = len(query_terms.intersection(re.findall(r"[a-z]{4,}", searchable)))
        score = overlap * 10 + (3 if paper.get("abstract") else 0) + (2 if paper.get("venue") else 0)
        ranked.append({**paper, "relevance_score": score, "metadata_complete": True})
    ranked.sort(key=lambda item: (item["relevance_score"], item.get("citation_count", 0)), reverse=True)
    # A search result is already topical; retain it when metadata is complete,
    # while limiting the prompt-sized working set to the strongest 25 records.
    ranked = ranked[:25]
    for index, paper in enumerate(ranked, 1):
        paper["source_id"] = f"SRC{index}"
    references = [{key: paper.get(key) for key in (
        "source_id", "title", "authors", "year", "doi", "url", "venue", "volume", "issue", "pages",
    )} for paper in ranked]
    print(f"Selected {len(ranked)} complete, ranked sources for analysis.")
    return {"literature_results": ranked, "references": references}


class LiteratureAnalysis(BaseModel):
    source_id: str = ""
    problem: str = "Not reported in the supplied abstract."
    methodology: str = "Not reported in the supplied abstract."
    dataset: str = "Not reported in the supplied abstract."
    models_or_algorithms: str = "Not reported in the supplied abstract."
    findings: str = "Not reported in the supplied abstract."
    limitations: str = "Not reported in the supplied abstract."
    relevance_to_project: str = ""


class LiteratureAnalysisReport(BaseModel):
    papers: list[LiteratureAnalysis] = Field(default_factory=list)


literature_analysis_llm = structured(LiteratureAnalysisReport)

LITERATURE_ANALYSIS_PROMPT = NO_INVENTION + """
You are the Paper Analysis Agent. Analyse each supplied retrieved paper using ONLY
its returned metadata and abstract. Produce one row per source_id for a literature
matrix: problem, methodology, dataset, models/algorithms, findings, limitations,
and relevance to the proposed project. If an item is absent from the abstract,
write exactly "Not reported in the supplied abstract." Never infer it from a title.
"""


def literature_analysis_node(state: ResearchState) -> dict:
    print("\n--- [Stage 6] Paper Analysis ---")
    report = run_agent(
        literature_analysis_llm,
        LITERATURE_ANALYSIS_PROMPT,
        {"PROJECT": state.get("project_profile", {}), "RANKED PAPERS": literature_brief(state, 1200)},
    )
    valid_ids = {paper.get("source_id") for paper in state.get("literature_results", [])}
    matrix = [row for row in report.get("papers", []) if row.get("source_id") in valid_ids]
    return {"literature_matrix": matrix}


class LiteratureSynthesis(BaseModel):
    themes: list[str] = Field(default_factory=list)
    comparative_insights: list[str] = Field(default_factory=list)
    trends: list[str] = Field(default_factory=list)
    evidence_based_limitations: list[str] = Field(default_factory=list)
    writing_guidance: list[str] = Field(default_factory=list)


literature_synthesis_llm = structured(LiteratureSynthesis)

LITERATURE_SYNTHESIS_PROMPT = NO_INVENTION + """
You are the Literature Synthesis Agent. Use the literature matrix to prepare a
thematic, comparative literature review plan. Group sources by meaningful themes,
identify evidence-based trends and limitations, and give writing guidance that
uses source_id citations. Do not write paper-by-paper summaries and do not claim
findings that are absent from the matrix.
"""


def literature_synthesis_node(state: ResearchState) -> dict:
    print("\n--- [Stage 7] Literature Synthesis ---")
    synthesis = run_agent(
        literature_synthesis_llm,
        LITERATURE_SYNTHESIS_PROMPT,
        {"PROJECT": state.get("project_profile", {}), "LITERATURE MATRIX": state.get("literature_matrix", [])},
    )
    return {"literature_synthesis": synthesis}



# 6. RESEARCH GAP



class ResearchGap(BaseModel):
    overall_gap: str = ""
    identified_gaps: list[str] = Field(default_factory=list)
    gap_evidence: list[str] = Field(default_factory=list)
    existing_approaches: list[str] = Field(default_factory=list)
    limitations_of_existing_work: list[str] = Field(default_factory=list)
    project_alignment: str = ""
    potential_contribution: str = ""
    confidence: str = ""


research_gap_llm = structured(ResearchGap)

RESEARCH_GAP_PROMPT = NO_INVENTION + """
You are the Research Gap Agent of PaperForge AI.
Identify research gaps by analysing the research problem, objectives, plan and the
retrieved literature (each paper has a source_id).
- Do not claim a gap unless the retrieved literature provides evidence for it.
- A missing search result is not proof that no research exists.
- Distinguish: (a) demonstrated gaps, (b) possible gaps needing investigation,
  (c) things that cannot currently be established.
- Cite evidence by source_id where possible.
- Decide whether the user's project addresses any gap; if evidence is
  insufficient, say so explicitly.
- Do not write the literature review or design a methodology.
"""


def research_gap_node(state: ResearchState) -> dict:
    print("\n--- [Stage 5] Research Gap ---")
    gap = run_agent(
        research_gap_llm,
        RESEARCH_GAP_PROMPT,
        {
            "RESEARCH PROBLEM": state.get("research_problem", ""),
            "RESEARCH OBJECTIVES": state.get("research_objectives", []),
            "RESEARCH PLAN": state.get("research_plan", {}),
            "RETRIEVED LITERATURE": literature_brief(state, 600),
            "LITERATURE MATRIX": state.get("literature_matrix", []),
            "LITERATURE SYNTHESIS": state.get("literature_synthesis", {}),
        },
    )
    return {"research_gap": gap}



# 6. PAPER OUTLINE



class OutlineSection(BaseModel):
    section_number: str = ""
    title: str
    purpose: str = ""
    key_points: list[str] = Field(default_factory=list)
    evidence_needed: list[str] = Field(default_factory=list)
    related_sources: list[str] = Field(default_factory=list)


class PaperOutline(BaseModel):
    title_candidates: list[str] = Field(default_factory=list)
    research_problem: str = ""
    research_objectives: list[str] = Field(default_factory=list)
    research_gap: str = ""
    contribution_summary: list[str] = Field(default_factory=list)
    sections: list[OutlineSection] = Field(default_factory=list)
    writing_notes: list[str] = Field(default_factory=list)


paper_outline_llm = structured(PaperOutline)

PAPER_OUTLINE_PROMPT = NO_INVENTION + """
You are the Paper Outline Agent of PaperForge AI.
Create a detailed IEEE-style paper outline based ONLY on the project profile,
research plan, research gap and retrieved literature. Do NOT write the sections.

Produce: several title candidates, research problem, objectives, research gap,
contribution summary, sections (with purpose, key_points, evidence_needed and
related_sources as source_ids), and writing notes.

Use EXACTLY these section titles, in this order (no numbering, no References
section, no separate Experimental Setup section, no separate Discussion
section):
Introduction, Literature Review, Methodology, Result and Discussion, Conclusion,
Future Scope.

- The gap must be reflected in Introduction, Literature Review and contributions.
- Methodology only from the described project methodology.
- Result and Discussion: state what must be reported (including any dataset,
  setup, hardware/software or evaluation-protocol details needed as context),
  then how those results should be interpreted against objectives, gap and
  literature; never fabricate results.
- If information is missing for a section, say the user must supply it.
"""


def paper_outline_node(state: ResearchState) -> dict:
    print("\n--- [Stage 6] Paper Outline ---")
    outline = run_agent(
        paper_outline_llm,
        PAPER_OUTLINE_PROMPT,
        {
            "PROJECT PROFILE": state.get("project_profile", {}),
            "RESEARCH PLAN": state.get("research_plan", {}),
            "RESEARCH GAP": state.get("research_gap", {}),
            "RETRIEVED LITERATURE": literature_brief(state, 300),
            "LITERATURE MATRIX": state.get("literature_matrix", []),
            "LITERATURE SYNTHESIS": state.get("literature_synthesis", {}),
            "HUMAN FEEDBACK ON PREVIOUS OUTLINE": state.get("approval_feedback", ""),
        },
    )
    titles = outline["title_candidates"]
    return {
        "paper_outline": outline,
        "title_candidates": titles,
        "selected_title": titles[0] if titles else "",
        "research_contributions": outline["contribution_summary"],
        # a new outline invalidates earlier writing progress
        "completed_sections": [],
    }



# 7. HUMAN APPROVAL



def human_approval_node(state: ResearchState) -> dict:
    print("\n--- [Human Approval] ---")
    user_response = interrupt(
        {
            "type": "paper_outline_approval",
            "message": "Please review the generated paper outline.",
            "title_candidates": state.get("title_candidates", []),
            "paper_outline": state.get("paper_outline", {}),
            "instructions": "Approve to continue to the Writing Agent, or reject with feedback.",
        }
    )
    status = str(user_response.get("status", "rejected")).strip().lower()
    update = {
        "approval_status": status,
        "approval_feedback": user_response.get("feedback", ""),
    }
    chosen = user_response.get("selected_title")
    if chosen:
        update["selected_title"] = chosen
    return update


def route_after_human_approval(state: ResearchState) -> str:
    if state.get("approval_status", "").strip().lower() == "approved":
        return "writer"
    return "research_planner"



# 8. WRITER



class WrittenSection(BaseModel):
    section_title: str = ""
    content: str = ""
    claims: list[str] = Field(default_factory=list)
    citations_used: list[str] = Field(default_factory=list)
    citations_needed: list[str] = Field(default_factory=list)
    unsupported_or_missing: list[str] = Field(default_factory=list)
    index_terms: list[str] = Field(default_factory=list)  # only used for the Abstract


writer_llm = structured(WrittenSection)

WRITER_PROMPT = NO_INVENTION + """
You are the Writing Agent of PaperForge AI. Write ONE section of an IEEE-style
research paper. Do NOT do final IEEE layout (a later step does that).

SOURCE PRIORITY: 1) uploaded materials / actual project evidence, 2) confirmed
project facts, 3) approved plan, 4) approved research gap, 5) retrieved literature.
The approved outline is authoritative: follow the section purpose and key points,
do not add objectives, contributions or change the research gap.

STYLE: formal, precise, objective, concise academic prose. No first person, no
promotional words ("revolutionary", "state-of-the-art" etc. unless evidenced).
Separate established findings, project observations, literature claims,
interpretations, limitations and future work. Never turn a possible gap into a
confirmed gap and never claim novelty merely because a paper was not retrieved.

CITATIONS: cite ONLY retrieved papers using their source_id in square brackets,
e.g. [SRC3] or [SRC1, SRC4]. Never invent citations, numbers, authors, DOIs.

SECTION GUIDES
- Introduction: context -> problem -> importance -> existing approaches ->
  limitations -> gap -> objective -> questions (if any) -> approach -> contributions.
- Literature Review: thematic and comparative (not paper-by-paper); limitations,
  datasets, evaluation; connect to the gap; only retrieved papers.
- Methodology: only what the project actually describes (architecture, data,
  preprocessing, training, setup, metrics, baselines). Do not add "standard"
  methods that were not stated.
- Result and Discussion: first report ONLY actual results provided (you may
  briefly note supplied datasets, splits, hardware, software, hyperparameters,
  baselines and evaluation protocol as context - explicitly say what was not
  provided instead of assuming it), then interpret those actual results against
  objectives, gap and literature; clearly mark speculation; state limitations.
  If no results exist, write only what is supported and list the missing result
  information in unsupported_or_missing.
- Conclusion: problem, approach, findings, contributions, limitations, future work;
  introduce nothing new.
- Future Scope: evidence-based next steps and stated limitations only; do not
  present planned work as completed work.
- Abstract (written last): background, objective, method, key ACTUAL results,
  contribution, conclusion - consistent with the manuscript. Also return
  index_terms (concise technical keywords from the paper only).

Keep terminology, dataset names, model names, metrics and numbers consistent with
previously written sections. Return only the section content plus schema metadata.
"""


def _pending_sections(state: ResearchState) -> list[dict]:
    done = set(state.get("completed_sections", []))
    pending, seen = [], set()
    for s in state.get("paper_outline", {}).get("sections", []):
        key = section_key(s.get("title", ""))
        if key and key != "abstract" and key not in done and key not in seen:
            pending.append({**s, "_key": key})
            seen.add(key)
    return pending


def writer_node(state: ResearchState) -> dict:
    print("\n--- [Stage 7] Writer ---")
    done = list(state.get("completed_sections", []))
    pending = _pending_sections(state)
    written = {k: state[k] for k in SECTION_KEYS if state.get(k)}

    common = {
        "APPROVED PAPER OUTLINE": state.get("paper_outline", {}),
        "PROJECT PROFILE": state.get("project_profile", {}),
        "ACTUAL RESULTS (verified)": state.get("actual_results", {}),
        "RESEARCH PLAN": state.get("research_plan", {}),
        "RESEARCH GAP": state.get("research_gap", {}),
        "RETRIEVED LITERATURE (cite by source_id)": literature_brief(state, 400),
        "LITERATURE MATRIX": state.get("literature_matrix", []),
        "LITERATURE SYNTHESIS": state.get("literature_synthesis", {}),
        "PREVIOUSLY WRITTEN SECTIONS": written,
    }

    if pending:
        section = pending[0]
        key = section.pop("_key")
        print(f"Writing section: {section.get('title')}")
        out = run_agent(writer_llm, WRITER_PROMPT, {**common, "CURRENT SECTION TO WRITE": section})
        return {key: out["content"], "completed_sections": done + [key]}

    if "abstract" not in done:
        print("Writing abstract and index terms")
        out = run_agent(
            writer_llm,
            WRITER_PROMPT,
            {**common, "CURRENT SECTION TO WRITE": {"title": "Abstract", "purpose": "Summarise the completed paper."}},
        )
        return {
            "abstract": out["content"],
            "index_terms": out["index_terms"],
            "completed_sections": done + ["abstract"],
        }

    return {}


def route_after_writer(state: ResearchState) -> str:
    if _pending_sections(state) or "abstract" not in state.get("completed_sections", []):
        return "writer"
    return "citation_verification"



# 9. CITATION VERIFICATION



class CitationVerificationResult(BaseModel):
    citation_id: str = ""
    citation_text: str = ""
    cited_claim: str = ""
    source_title: str = ""
    source_authors: list[str] = Field(default_factory=list)
    source_year: str = ""
    source_doi: str = ""
    source_url: str = ""
    source_found: bool = False
    claim_supported: bool = False
    support_level: str = ""  # supported | partially_supported | unsupported | unclear
    evidence_from_source: str = ""
    issue: str = ""
    recommendation: str = ""
    section: str = ""


class CitationVerificationReport(BaseModel):
    overall_status: str = ""  # pass | needs_revision | insufficient_evidence
    total_citations: int = 0
    verified_citations: int = 0
    unsupported_citations: int = 0
    partial_citations: int = 0
    unclear_citations: int = 0
    missing_sources: int = 0
    results: list[CitationVerificationResult] = Field(default_factory=list)
    summary: str = ""
    critical_issues: list[str] = Field(default_factory=list)


citation_verification_llm = structured(CitationVerificationReport)

CITATION_VERIFICATION_PROMPT = NO_INVENTION + """
You are the Citation Verification Agent of PaperForge AI. VERIFICATION ONLY:
do not rewrite the manuscript or generate replacement citations.

For every citation ([SRCn] identifiers) in the manuscript:
1. identify the cited claim and its section
2. find the source in the supplied literature / citation records
3. check the source exists and metadata is consistent
4. decide whether the source actually supports the claim
   (relevance is NOT support; a claim stronger than the evidence is partial)
5. flag citations attached to unrelated claims, one citation covering several
   claims, and important externally verifiable claims lacking a citation.

support_level must be exactly one of: supported, partially_supported,
unsupported, unclear. Use only supplied evidence, never general knowledge.
If only an abstract is available, judge only on the abstract; if insufficient,
use "unclear".

overall_status: pass ONLY if all citations are adequately supported;
needs_revision if any citation problem needs correction; insufficient_evidence if
source information is too thin to verify important citations.
Give a recommendation for the Revision Agent for each problem.

Never mark a citation as supported merely because its title is topically similar.
When metadata is incomplete (missing title, author list, year, and both DOI/URL),
flag the source as unverifiable. Flag an externally-derived claim without a
nearby citation as a missing-citation issue and name its section.
"""

_CITE_RE = re.compile(r"\[(SRC\d+(?:\s*,\s*SRC\d+)*)\]")


def build_citation_records(state: ResearchState) -> list[dict]:
    """Deterministically map each source_id to the sections that cite it."""
    cited: dict[str, list[str]] = {}
    for title, key in [("Abstract", "abstract")] + SECTION_ORDER:
        for m in _CITE_RE.finditer(state.get(key, "") or ""):
            for sid in re.split(r"\s*,\s*", m.group(1)):
                cited.setdefault(sid, [])
                if title not in cited[sid]:
                    cited[sid].append(title)
    refs = {r.get("source_id"): r for r in state.get("references", [])}
    records = []
    for sid, sections in cited.items():
        r = refs.get(sid, {})
        records.append(
            {
                "citation_id": sid,
                "source_title": r.get("title", "NOT FOUND IN RETRIEVED LITERATURE"),
                "source_authors": r.get("authors", []),
                "source_doi": r.get("doi", ""),
                "source_url": r.get("url", ""),
                "source_year": r.get("year", ""),
                "source_venue": r.get("venue", ""),
                "sections": sections,
            }
        )
    return records


def citation_verification_node(state: ResearchState) -> dict:
    print("\n--- [Stage 10] Citation Verification ---")
    records = build_citation_records(state)
    report = run_agent(
        citation_verification_llm,
        CITATION_VERIFICATION_PROMPT,
        {
            "MANUSCRIPT": build_manuscript(state),
            "CITATION RECORDS": records,
            "RETRIEVED LITERATURE": literature_brief(state, 700),
            "PREVIOUS CITATION FEEDBACK": state.get("citation_feedback", ""),
        },
    )
    issues = []
    for r in report.get("results", []):
        if (not r.get("source_found")) or r.get("support_level") in {
            "unsupported", "partially_supported", "unclear",
        }:
            issues.append(
                {
                    "citation_id": r.get("citation_id", ""),
                    "claim": r.get("cited_claim", ""),
                    "section": r.get("section", ""),
                    "source_title": r.get("source_title", ""),
                    "issue": r.get("issue", ""),
                    "support_level": r.get("support_level", ""),
                    "recommendation": r.get("recommendation", ""),
                }
            )
    cited_ids = {record.get("citation_id") for record in records if record.get("citation_id")}
    reference_by_id = {ref.get("source_id"): ref for ref in state.get("references", []) if isinstance(ref, dict)}
    manuscript = build_manuscript(state)
    if state.get("literature_results") and not cited_ids:
        issues.append({
            "citation_id": "",
            "claim": "The manuscript contains retrieved literature but no in-text citations.",
            "section": "Manuscript",
            "source_title": "",
            "issue": "No retrieved source is cited in the paper.",
            "support_level": "unsupported",
            "recommendation": "Add only relevant [SRCn] citations next to literature-derived claims.",
        })
    for source_id in cited_ids:
        ref = reference_by_id.get(source_id, {})
        has_locator = bool(ref.get("doi") or ref.get("url"))
        if not ref.get("title") or not ref.get("authors") or not ref.get("year") or not has_locator:
            issues.append({
                "citation_id": source_id,
                "claim": "Citation source metadata",
                "section": ", ".join(next((r.get("sections", []) for r in records if r.get("citation_id") == source_id), [])),
                "source_title": ref.get("title", ""),
                "issue": "Citation metadata is incomplete; the source cannot be reliably identified.",
                "support_level": "unclear",
                "recommendation": "Replace the citation with a retrieved source that has title, authors, year, and DOI or URL.",
            })
    malformed = [match.group(0) for match in re.finditer(r"\[SRC[^\]]*\]", manuscript) if not _CITE_RE.fullmatch(match.group(0))]
    for marker in sorted(set(malformed)):
        issues.append({
            "citation_id": marker,
            "claim": "Citation marker syntax",
            "section": "Manuscript",
            "source_title": "",
            "issue": "Citation marker is not in the required [SRC1] or [SRC1, SRC2] format.",
            "support_level": "unsupported",
            "recommendation": "Use a comma-separated [SRCn] marker that corresponds to retrieved literature.",
        })
    entry = {
        "verification_type": "citation",
        "status": report.get("overall_status", ""),
        "total_citations": report.get("total_citations", 0),
        "unsupported_citations": report.get("unsupported_citations", 0) + len(issues),
        "issues": issues,
    }
    return {
        "citation_records": records,
        "citation_verification_results": report.get("results", []),
        "citation_feedback": report.get("summary", ""),
        "citation_issues": issues,
        "verification_results": [entry],
    }



# 10. CLAIM VERIFICATION



class ClaimVerificationResult(BaseModel):
    claim_id: str = ""
    claim_text: str = ""
    section: str = ""
    claim_type: str = ""
    evidence_source: str = ""
    evidence_available: bool = False
    claim_supported: bool = False
    support_level: str = ""
    evidence_excerpt: str = ""
    issue: str = ""
    recommendation: str = ""


class ClaimVerificationReport(BaseModel):
    overall_status: str = ""
    total_claims: int = 0
    verified_claims: int = 0
    supported_claims: int = 0
    partial_claims: int = 0
    unsupported_claims: int = 0
    unclear_claims: int = 0
    results: list[ClaimVerificationResult] = Field(default_factory=list)
    summary: str = ""
    critical_issues: list[str] = Field(default_factory=list)


claim_verification_llm = structured(ClaimVerificationReport)

CLAIM_VERIFICATION_PROMPT = NO_INVENTION + """
You are the Claim Verification Agent of PaperForge AI. VERIFICATION ONLY: do not
rewrite the manuscript.

Identify the important factual and scientific claims and check each against the
supplied evidence (project profile, actual results, literature, citation
verification). Never use general model knowledge as evidence.

claim_type: project_fact | experimental_result | numerical_result |
methodological_claim | literature_claim | interpretation | conclusion | other.
support_level: supported | partially_supported | unsupported | unclear.
- Numerical claims (accuracy, F1, dataset sizes, splits, percentages,
  improvements) must match supplied evidence exactly.
- Claimed experiments/methods must appear in the project information.
- Watch for conclusions stronger than results, unsupported causal claims,
  over-generalisation, superiority claims without comparison.
- Missing evidence -> unclear; evidence contradicting the claim -> unsupported.
Give a recommendation for the Revision Agent for each problem.

overall_status: pass | needs_revision | insufficient_evidence.
"""


def claim_verification_node(state: ResearchState) -> dict:
    print("\n--- [Stage 11] Claim Verification ---")
    report = run_agent(
        claim_verification_llm,
        CLAIM_VERIFICATION_PROMPT,
        {
            "MANUSCRIPT": build_manuscript(state),
            "PROJECT PROFILE": state.get("project_profile", {}),
            "ACTUAL PROJECT RESULTS": state.get("actual_results", {}),
            "RETRIEVED LITERATURE": literature_brief(state, 500),
            "CITATION ISSUES": state.get("citation_issues", []),
            "PREVIOUS CLAIM FEEDBACK": state.get("claim_feedback", ""),
        },
    )
    issues = []
    for r in report.get("results", []):
        if (not r.get("evidence_available")) or r.get("support_level") in {
            "partially_supported", "unsupported", "unclear",
        }:
            issues.append(
                {
                    "claim_id": r.get("claim_id", ""),
                    "claim": r.get("claim_text", ""),
                    "section": r.get("section", ""),
                    "claim_type": r.get("claim_type", ""),
                    "evidence_source": r.get("evidence_source", ""),
                    "support_level": r.get("support_level", ""),
                    "issue": r.get("issue", ""),
                    "recommendation": r.get("recommendation", ""),
                }
            )
    entry = {
        "verification_type": "claim",
        "status": report.get("overall_status", ""),
        "total_claims": report.get("total_claims", 0),
        "unsupported_claims": report.get("unsupported_claims", 0),
        "issues": issues,
    }
    return {
        "claim_feedback": report.get("summary", ""),
        "claim_issues": issues,
        "verification_results": [entry],
    }



# 11. CONSISTENCY VERIFICATION



class ConsistencyVerificationResult(BaseModel):
    consistency_id: str = ""
    issue_type: str = ""
    sections_involved: list[str] = Field(default_factory=list)
    statement_a: str = ""
    statement_b: str = ""
    consistent: bool = True
    severity: str = ""  # low | medium | high | critical
    explanation: str = ""
    evidence: str = ""
    recommendation: str = ""


class ConsistencyVerificationReport(BaseModel):
    overall_status: str = ""
    total_checks: int = 0
    consistent_checks: int = 0
    inconsistent_checks: int = 0
    high_severity_issues: int = 0
    critical_issues: int = 0
    results: list[ConsistencyVerificationResult] = Field(default_factory=list)
    summary: str = ""
    critical_issues_list: list[str] = Field(default_factory=list)


consistency_verification_llm = structured(ConsistencyVerificationReport)

CONSISTENCY_VERIFICATION_PROMPT = NO_INVENTION + """
You are the Consistency Verification Agent of PaperForge AI. VERIFICATION ONLY:
do not rewrite the manuscript.

Check internal consistency from Abstract to Conclusion: numerical values
(accuracy, F1, dataset size, splits, counts), methodology vs the Result and
Discussion section, dataset description, evaluation metrics, objectives,
contributions, terminology (only when different names could mean different
systems), the Result and Discussion section vs Conclusion, Abstract vs body,
and citation usage.

Only report a conflict when the statements cannot both reasonably be true or the
difference causes real ambiguity; never flag harmless wording differences.
If a conflict cannot be established from the evidence, do not declare it.

For each inconsistency: issue_type, sections_involved, statement_a, statement_b,
consistent=false, severity (low|medium|high|critical), explanation, evidence,
recommendation. Only list inconsistencies in results.
overall_status: pass | needs_revision | insufficient_evidence.
"""


def consistency_verification_node(state: ResearchState) -> dict:
    print("\n--- [Stage 12] Consistency Verification ---")
    report = run_agent(
        consistency_verification_llm,
        CONSISTENCY_VERIFICATION_PROMPT,
        {
            "MANUSCRIPT": build_manuscript(state),
            "PROJECT PROFILE": state.get("project_profile", {}),
            "ACTUAL PROJECT RESULTS": state.get("actual_results", {}),
            "CLAIM ISSUES": state.get("claim_issues", []),
            "PREVIOUS CONSISTENCY FEEDBACK": state.get("consistency_feedback", ""),
        },
    )
    issues = [
        {
            "consistency_id": r.get("consistency_id", ""),
            "issue_type": r.get("issue_type", ""),
            "sections_involved": r.get("sections_involved", []),
            "statement_a": r.get("statement_a", ""),
            "statement_b": r.get("statement_b", ""),
            "severity": r.get("severity", ""),
            "explanation": r.get("explanation", ""),
            "evidence": r.get("evidence", ""),
            "recommendation": r.get("recommendation", ""),
        }
        for r in report.get("results", [])
        if not r.get("consistent", True)
    ]
    entry = {
        "verification_type": "consistency",
        "status": report.get("overall_status", ""),
        "inconsistent_checks": report.get("inconsistent_checks", 0),
        "issues": issues,
    }
    return {
        "consistency_feedback": report.get("summary", ""),
        "consistency_issues": issues,
        "verification_results": [entry],
    }



# 12. ACADEMIC REVIEW



class ReviewIssue(BaseModel):
    issue_id: str = ""
    issue_type: str = ""
    severity: str = ""
    section: str = ""
    description: str = ""
    evidence: str = ""
    recommendation: str = ""


class ReviewReport(BaseModel):
    overall_status: str = ""
    issues_found: bool = False
    total_issues: int = 0
    high_severity_issues: int = 0
    critical_issues: int = 0
    issues: list[ReviewIssue] = Field(default_factory=list)
    summary: str = ""
    revision_required: bool = False


review_llm = structured(ReviewReport)

REVIEW_PROMPT = NO_INVENTION + """
You are the Academic Review Agent of PaperForge AI. REVIEW / QUALITY CONTROL
only; do not rewrite the manuscript.

Review the manuscript together with the citation, claim and consistency
verification outputs. Decide if meaningful issues must be fixed before
formatting: citation problems, unsupported or partially supported claims,
numerical/methodology/dataset/results inconsistencies, structural gaps, and
writing problems that materially hurt clarity.

Take high/critical verification issues seriously, but do not blindly accept every
flagged item; judge with the evidence.

Severity: low | medium | high | critical. For each issue give issue_id,
issue_type, severity, section (e.g. "Result and Discussion"), description,
evidence, recommendation.

No meaningful issues -> issues_found=false, revision_required=false,
overall_status="pass". Otherwise issues_found=true, revision_required=true,
overall_status="needs_revision". Use "insufficient_evidence" only when missing
evidence prevents reliable review.
"""


def review_node(state: ResearchState) -> dict:
    print("\n--- [Stage 13] Academic Review ---")
    report = run_agent(
        review_llm,
        REVIEW_PROMPT,
        {
            "CURRENT MANUSCRIPT": build_manuscript(state),
            "CITATION ISSUES": state.get("citation_issues", []),
            "CLAIM ISSUES": state.get("claim_issues", []),
            "CONSISTENCY ISSUES": state.get("consistency_issues", []),
            "PREVIOUS REVIEW FEEDBACK": state.get("review_feedback", ""),
        },
    )
    issues = report.get("issues", [])
    revision_required = bool(report.get("revision_required")) or len(issues) > 0
    return {
        "review_feedback": report.get("summary", ""),
        "review_issues": issues,
        "is_approved": not revision_required,
        "verification_results": [
            {
                "verification_type": "review",
                "status": report.get("overall_status", ""),
                "revision_required": revision_required,
                "total_issues": report.get("total_issues", len(issues)),
                "issues": issues,
            }
        ],
    }


def route_after_review(state: ResearchState) -> str:
    """
    Revision if the reviewer did not approve, or a high/critical consistency issue
    remains. Capped at MAX_REVISIONS so the loop always terminates.
    """
    if state.get("revision_count", 0) >= MAX_REVISIONS:
        print(f"Revision cap ({MAX_REVISIONS}) reached - proceeding to formatting.")
        return "formatting"
    serious = any(
        i.get("severity") in {"high", "critical"}
        for i in state.get("consistency_issues", [])
        if isinstance(i, dict)
    )
    if not state.get("is_approved", False) or serious:
        return "revision"
    return "formatting"



# 13. REVISION



class RevisionResult(BaseModel):
    section_title: str = ""
    revised_content: str = ""
    changes_made: list[str] = Field(default_factory=list)
    unresolved_issues: list[str] = Field(default_factory=list)


revision_llm = structured(RevisionResult)

REVISION_PROMPT = NO_INVENTION + """
You are the Revision Agent of PaperForge AI. Revise ONE section of the paper using
the supplied verification and review feedback.
- Fix the identified problems while preserving verified information.
- Do not change verified experimental results and do not redesign the research.
- Keep citations in [SRCn] form; only cite retrieved sources.
- If an issue cannot be resolved because evidence is missing, keep the supported
  content, soften/remove unsupported claims, and list it in unresolved_issues.
- Maintain IEEE-style academic language and consistent terminology.
Return section_title, revised_content, changes_made, unresolved_issues.
"""


def _build_revision_issues(state: ResearchState) -> list[dict]:
    issues = []
    for issue_type, key in [
        ("review", "review_issues"),
        ("citation", "citation_issues"),
        ("claim", "claim_issues"),
        ("consistency", "consistency_issues"),
    ]:
        for issue in state.get(key, []):
            item = dict(issue) if isinstance(issue, dict) else {"issue": str(issue)}
            item["issue_type"] = issue_type
            issues.append(item)
    for issue in state.get("final_validation_issues", []):
        issues.append({"issue_type": "final_validation", "issue": str(issue)})
    return issues


def _issue_sections(issue: dict) -> list[str]:
    names = [issue.get("section", "")] + list(issue.get("sections_involved", []))
    return [k for k in (section_key(str(n)) for n in names if n) if k]


def _get_sections_to_revise(state: ResearchState, issues: list[dict]) -> list[str]:
    sections: list[str] = []
    for issue in issues:
        for k in _issue_sections(issue):
            if k not in sections:
                sections.append(k)
        if issue.get("issue_type") == "final_validation":
            text = str(issue.get("issue", "")).lower()
            for alias, k in _KEY_MAP.items():
                if alias in text and k not in sections:
                    sections.append(k)
    if not sections:  # reviewers named no section -> revise everything that exists
        sections = [k for k in SECTION_KEYS + ["abstract"] if state.get(k, "").strip()]
    return sections


def revision_node(state: ResearchState) -> dict:
    print("\n--- [Stage 14] Revision ---")
    count = state.get("revision_count", 0)
    issues = _build_revision_issues(state)
    targets = _get_sections_to_revise(state, issues)
    manuscript = build_manuscript(state)

    revised: dict[str, str] = {}
    changes: list[str] = []
    unresolved: list[str] = []

    for key in targets:
        current = state.get(key, "")
        if not current.strip():
            continue
        relevant = [i for i in issues if not _issue_sections(i) or key in _issue_sections(i)]
        try:
            out = run_agent(
                revision_llm,
                REVISION_PROMPT,
                {
                    "SECTION TO REVISE": key,
                    "CURRENT SECTION": current,
                    "ISSUES TO FIX": relevant,
                    "PROJECT PROFILE": state.get("project_profile", {}),
                    "CONFIRMED FACTS": state.get("confirmed_facts", []),
                    "ACTUAL EXPERIMENTAL RESULTS": state.get("actual_results", {}),
                    "RESEARCH GAP": state.get("research_gap", {}),
                    "RETRIEVED LITERATURE": literature_brief(state, 400),
                    "FULL CURRENT MANUSCRIPT": manuscript,
                },
            )
            text = out.get("revised_content", "").strip()
            if text:
                revised[key] = text
            changes += out.get("changes_made", [])
            unresolved += out.get("unresolved_issues", [])
            print(f"Revised section: {key}")
        except Exception as exc:  # noqa: BLE001
            print(f"Revision failed for {key}: {exc}")
            unresolved.append(f"Revision failed for {key}: {exc}")

    summary = f"Revision cycle {count + 1}. Sections revised: {', '.join(revised) or 'none'}."
    if changes:
        summary += " Changes: " + "; ".join(changes)
    if unresolved:
        summary += " Unresolved: " + "; ".join(unresolved)

    update = {
        "revision_count": count + 1,
        "current_stage": "revision",
        "is_approved": False,
        "review_feedback": summary,
        "citation_feedback": "",
        "claim_feedback": "",
        "consistency_feedback": "",
        "review_issues": [],
        "citation_issues": [],
        "claim_issues": [],
        "consistency_issues": [],
        "final_validation_status": "",
        "final_validation_issues": [],
        "final_validation_warnings": [],
    }
    update.update(revised)
    return update



# EXPORT PREPARATION  (numeric IEEE citations: [SRC3] -> [1])



def _ref_string(ref: Any) -> str:
    if not isinstance(ref, dict):
        return str(ref)
    authors = ref.get("authors", "")
    if isinstance(authors, list):
        def ieee_name(name: Any) -> str:
            words = str(name).replace(",", " ").split()
            if len(words) < 2:
                return str(name)
            return " ".join(f"{word[0]}." for word in words[:-1]) + f" {words[-1]}"
        rendered = [ieee_name(author) for author in authors if author]
        if len(rendered) == 1:
            authors = rendered[0]
        elif len(rendered) == 2:
            authors = f"{rendered[0]} and {rendered[1]}"
        elif rendered:
            authors = ", ".join(rendered[:-1]) + f", and {rendered[-1]}"
        else:
            authors = ""
    parts: list[str] = []
    if authors:
        parts.append(str(authors))
    if ref.get("title"):
        parts.append(f'"{ref["title"]}"')
    if ref.get("venue"):
        parts.append(str(ref["venue"]))
    volume_issue = []
    if ref.get("volume"):
        volume_issue.append(f"vol. {ref['volume']}")
    if ref.get("issue"):
        volume_issue.append(f"no. {ref['issue']}")
    if ref.get("pages"):
        volume_issue.append(f"pp. {ref['pages']}")
    if volume_issue:
        parts.append(", ".join(volume_issue))
    if ref.get("year"):
        parts.append(str(ref["year"]))
    if ref.get("doi"):
        doi = str(ref["doi"]).replace("https://doi.org/", "")
        parts.append(f"doi: {doi}")
    elif ref.get("url"):
        parts.append(f"[Online]. Available: {ref['url']}")
    return ", ".join(parts)


def prepare_export(state: ResearchState) -> dict:
    """Return title/abstract/sections/references with numeric citations applied."""
    refs = state.get("references", [])
    by_id = {r.get("source_id"): r for r in refs if isinstance(r, dict)}

    texts = {"abstract": state.get("abstract", "")}
    for _, key in SECTION_ORDER:
        texts[key] = state.get(key, "")

    order: list[str] = []
    for key in ["abstract"] + SECTION_KEYS:
        for m in _CITE_RE.finditer(texts[key] or ""):
            for sid in re.split(r"\s*,\s*", m.group(1)):
                if sid not in order:
                    order.append(sid)

    number = {sid: i + 1 for i, sid in enumerate(order)}

    def sub(text: str) -> str:
        return _CITE_RE.sub(
            lambda m: "[" + ", ".join(str(number[s]) for s in re.split(r"\s*,\s*", m.group(1))) + "]",
            text or "",
        )

    # IEEE references are only emitted when they are cited in the manuscript.
    # This prevents an uncited bibliography and exposes unresolved source IDs to
    # final validation instead of silently exporting placeholders.
    ref_strings = [
        _ref_string(by_id[sid]) if sid in by_id else f"Unknown source {sid}"
        for sid in order
    ]

    return {
        "title": state.get("selected_title", ""),
        "abstract": sub(texts["abstract"]),
        "index_terms": state.get("index_terms", []),
        "sections": [(t, sub(texts[k])) for t, k in SECTION_ORDER if texts[k]],
        "references": ref_strings,
        "missing_refs": [sid for sid in order if sid not in by_id],
    }


def render_markdown(state: ResearchState) -> str:
    d = prepare_export(state)
    parts = []
    if d["title"]:
        parts.append(f"# {d['title']}")
    if d["abstract"]:
        parts.append(f"## Abstract\n\n{d['abstract']}")
    if d["index_terms"]:
        parts.append("## Keywords\n\n" + ", ".join(d["index_terms"]))
    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
    for i, (title, text) in enumerate(d["sections"]):
        parts.append(f"## {roman[i]}. {title}\n\n{text}")
    parts.append("## References\n\n" + "\n".join(f"[{i}] {r}" for i, r in enumerate(d["references"], 1)))
    return "\n\n".join(parts)



# 14. FORMATTING



class FormattingResult(BaseModel):
    title: str = ""
    abstract: str = ""
    index_terms: list[str] = Field(default_factory=list)
    introduction: str = ""
    literature_review: str = ""
    methodology: str = ""
    result_and_discussion: str = ""
    conclusion: str = ""
    future_scope: str = ""
    formatting_notes: list[str] = Field(default_factory=list)


formatting_llm = structured(FormattingResult)

FORMATTING_PROMPT = NO_INVENTION + """
You are the Formatting Agent of PaperForge AI. FORMATTING ONLY.
Organise the verified manuscript into a clean IEEE-style structure (title,
abstract, Keywords, Introduction, Literature Review, Methodology, Result and
Discussion, Conclusion, Future Scope).
- Do NOT change scientific content, facts, numbers, dataset/model names, metrics.
- Keep citation markers such as [SRC3] EXACTLY as they are.
- Do not add claims, results, references, figures or tables.
- Use the selected title if provided; keep existing keywords.
- Only tidy paragraphs, whitespace and structure. If a section is fine, return it as is.
"""


def _prefer(new: str, old: str) -> str:
    """Use the formatted text unless it is empty or suspiciously shorter (content loss)."""
    if new and len(new) >= 0.7 * len(old or ""):
        return new
    return old or new or ""


def formatting_node(state: ResearchState) -> dict:
    print("\n--- [Stage 15] Formatting ---")
    formatted: dict = {}
    try:
        formatted = run_agent(
            formatting_llm,
            FORMATTING_PROMPT,
            {
                "VERIFIED MANUSCRIPT": build_manuscript(state),
                "SELECTED TITLE": state.get("selected_title", ""),
                "TITLE CANDIDATES": state.get("title_candidates", []),
                "INDEX TERMS": state.get("index_terms", []),
            },
        )
    except Exception as exc:  # noqa: BLE001
        print(f"Formatting LLM failed, keeping original text: {exc}")

    update: dict[str, Any] = {}
    for key in ["abstract"] + SECTION_KEYS:
        update[key] = _prefer(formatted.get(key, ""), state.get(key, ""))

    titles = state.get("title_candidates", [])
    update["selected_title"] = (
        state.get("selected_title") or formatted.get("title") or (titles[0] if titles else "")
    )
    update["index_terms"] = formatted.get("index_terms") or state.get("index_terms", [])

    merged = {**state, **update}
    update["final_manuscript"] = render_markdown(merged)
    update["verification_results"] = [
        {"verification_type": "formatting", "status": "completed", "notes": formatted.get("formatting_notes", [])}
    ]
    return update



# 15. FINAL VALIDATION



class FinalValidationResult(BaseModel):
    overall_status: str = ""  # pass | needs_revision | insufficient_evidence
    manuscript_present: bool = False
    citations_valid: bool = True
    claims_valid: bool = True
    consistency_valid: bool = True
    review_passed: bool = False
    issues: list[str] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)
    summary: str = ""


final_validation_llm = structured(FinalValidationResult)

FINAL_VALIDATION_PROMPT = NO_INVENTION + """
You are the Final Validation Agent of PaperForge AI. VALIDATION ONLY: do not
rewrite anything. Decide whether the manuscript is ready to export to DOCX/PDF/LaTeX.

Check: required content present (title, abstract, index terms, all major
sections, references when citations are used); no placeholder text (TBD, TODO,
INSERT HERE, PLACEHOLDER, XXX); references not obviously empty or disconnected
from citations; no unresolved CRITICAL verification issue.

Status: pass (ready), needs_revision (unresolved issue must be fixed),
insufficient_evidence (missing evidence prevents reliable validation).
List each concrete problem in issues (short strings; name the affected section
when relevant) and minor concerns in warnings.
"""

_PLACEHOLDER_RE = re.compile(r"\b(TBD|TODO|INSERT HERE|PLACEHOLDER|YOUR TEXT HERE|ADD FIGURE|ADD TABLE|XXX)\b")


def final_validation_node(state: ResearchState) -> dict:
    print("\n--- [Stage 16] Final Validation ---")
    llm_out = run_agent(
        final_validation_llm,
        FINAL_VALIDATION_PROMPT,
        {
            "FINAL MANUSCRIPT": state.get("final_manuscript", ""),
            "REFERENCES": state.get("references", []),
            "CITATION ISSUES": state.get("citation_issues", []),
            "CLAIM ISSUES": state.get("claim_issues", []),
            "CONSISTENCY ISSUES": state.get("consistency_issues", []),
            "REVIEW ISSUES": state.get("review_issues", []),
        },
        limit=MANUSCRIPT_CHARS,
    )

    # Deterministic checks
    det: list[str] = []
    if not state.get("final_manuscript", "").strip():
        det.append("Final manuscript is empty.")
    for label, key in [
        ("Title", "selected_title"),
        ("Abstract", "abstract"),
        ("Introduction", "introduction"),
        ("Methodology", "methodology"),
        ("Result and Discussion", "result_and_discussion"),
        ("Conclusion", "conclusion"),
        ("Future Scope", "future_scope"),
    ]:
        if not str(state.get(key, "")).strip():
            det.append(f"{label} is missing.")
    if _PLACEHOLDER_RE.search(state.get("final_manuscript", "")):
        det.append("Placeholder text remains in the manuscript.")

    # Re-check citation integrity after revision/formatting.  The workflow may
    # revise text after the dedicated citation agent has run, so export must
    # never rely on a stale verification result.
    current_records = build_citation_records(state)
    cited_ids = {record.get("citation_id") for record in current_records if record.get("citation_id")}
    reference_by_id = {ref.get("source_id"): ref for ref in state.get("references", []) if isinstance(ref, dict)}
    if state.get("literature_results") and not cited_ids:
        det.append("Retrieved literature is present but no in-text citations are used.")
    for source_id in cited_ids:
        ref = reference_by_id.get(source_id)
        if not ref:
            det.append(f"Citation {source_id} has no retrieved reference entry.")
        elif not ref.get("title") or not ref.get("authors") or not ref.get("year") or not (ref.get("doi") or ref.get("url")):
            det.append(f"Citation {source_id} has incomplete source metadata.")
    malformed = [m.group(0) for m in re.finditer(r"\[SRC[^\]]*\]", state.get("final_manuscript", "")) if not _CITE_RE.fullmatch(m.group(0))]
    if malformed:
        det.append("Malformed source citation marker(s): " + ", ".join(sorted(set(malformed))) + ".")

    warnings = list(llm_out.get("warnings", []))
    export = prepare_export(state)
    if export["missing_refs"]:
        det.append(f"Citations without reference entries: {', '.join(export['missing_refs'])}.")

    status_llm = llm_out.get("overall_status", "")
    if det:
        status = "needs_revision"
    elif status_llm == "pass":
        status = "pass"
    elif status_llm == "insufficient_evidence":
        status = "insufficient_evidence"
    else:
        status = "needs_revision"

    issues = det + list(llm_out.get("issues", []))
    return {
        "final_validation_status": status,
        "final_validation_issues": issues,
        "final_validation_warnings": warnings,
        "verification_results": [
            {
                "verification_type": "final_validation",
                "status": status,
                "issues": issues,
                "warnings": warnings,
                "summary": llm_out.get("summary", ""),
            }
        ],
    }


def route_after_final_validation(state: ResearchState) -> str:
    if state.get("final_validation_status", "").strip().lower() == "pass":
        return "document_export"
    if state.get("revision_count", 0) >= MAX_REVISIONS:
        print("Revision cap reached - exporting with unresolved validation issues.")
        return "document_export"
    return "revision"



# 16. DOCUMENT EXPORT



def export_docx(state: ResearchState, output_dir: str) -> str:
    d = prepare_export(state)
    path = os.path.join(output_dir, "PaperForge_Final.docx")
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(10)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    first = doc.sections[0]
    first.top_margin = first.bottom_margin = Inches(.75)
    first.left_margin = first.right_margin = Inches(.7)

    def para(text="", center=False, bold=False, size=10):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(text)
        r.bold, r.font.name, r.font.size = bold, "Times New Roman", Pt(size)
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        return p

    def columns(section, count):
        cols = section._sectPr.first_child_found_in("w:cols")
        if cols is None:
            cols = OxmlElement("w:cols")
            section._sectPr.append(cols)
        cols.set(qn("w:num"), str(count))
        cols.set(qn("w:space"), "360")

    if d["title"]:
        para(d["title"], center=True, bold=True, size=16)
    for author in state.get("paper_metadata", {}).get("authors", []):
        fields = [x.strip() for x in str(author).split("|")]
        para(fields[0], center=True)
        if len(fields) > 1:
            para("\n".join(fields[1:]), center=True, size=8)

    body = doc.add_section(WD_SECTION.CONTINUOUS)
    body.top_margin = body.bottom_margin = Inches(.75)
    body.left_margin = body.right_margin = Inches(.7)
    columns(body, 2)
    if d["abstract"]:
        para("Abstract— " + d["abstract"], bold=True)
    if d["index_terms"]:
        para("Keywords— " + ", ".join(d["index_terms"]), bold=True)

    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
    for i, (title, text) in enumerate(d["sections"]):
        para(f"{roman[i]}. {title.upper()}", center=True, bold=True)
        for text_para in re.split(r"\n\s*\n", text):
            if text_para.strip():
                para(text_para.strip())

    para("REFERENCES", center=True, bold=True)
    for i, r in enumerate(d["references"], 1):
        para(f"[{i}] {r}", size=8)

    doc.save(path)
    return path


def _pdf_escape(s: str) -> str:
    return str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def export_pdf(state: ResearchState, output_dir: str) -> str:
    d = prepare_export(state)
    path = os.path.join(output_dir, "PaperForge_Final.pdf")
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(path, pagesize=A4, rightMargin=50, leftMargin=50, topMargin=50, bottomMargin=50)
    story: list = []

    def block(text: str):
        story.append(Paragraph(_pdf_escape(text).replace("\n", "<br/>"), styles["BodyText"]))
        story.append(Spacer(1, 8))

    if d["title"]:
        story.append(Paragraph(_pdf_escape(d["title"]), styles["Title"]))
        story.append(Spacer(1, 12))
    if d["abstract"]:
        story.append(Paragraph("Abstract", styles["Heading1"]))
        block(d["abstract"])
    if d["index_terms"]:
        story.append(Paragraph("Keywords", styles["Heading1"]))
        block(", ".join(d["index_terms"]))

    roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
    for i, (title, text) in enumerate(d["sections"]):
        story.append(Paragraph(f"{roman[i]}. {_pdf_escape(title)}", styles["Heading1"]))
        block(text)

    story.append(Paragraph("References", styles["Heading1"]))
    for i, r in enumerate(d["references"], 1):
        block(f"[{i}] {r}")

    doc.build(story)
    return path


_TEX_SPECIALS = [
    ("\\", r"\textbackslash{}"), ("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"),
    ("_", r"\_"), ("{", r"\{"), ("}", r"\}"), ("~", r"\textasciitilde{}"), ("^", r"\textasciicircum{}"),
]


def _tex(s: str) -> str:
    out = str(s)
    # Escape backslash first via placeholder so later replacements don't touch it.
    out = out.replace("\\", "\x00")
    for ch, rep in _TEX_SPECIALS[1:]:
        out = out.replace(ch, rep)
    return out.replace("\x00", _TEX_SPECIALS[0][1])


def export_latex(state: ResearchState, output_dir: str) -> str:
    d = prepare_export(state)
    path = os.path.join(output_dir, "PaperForge_Final.tex")
    tex = [
        "\\documentclass[conference]{IEEEtran}",
        "\\usepackage{graphicx,amsmath,booktabs,url}",
        "\\begin{document}",
        f"\\title{{{_tex(d['title'])}}}",
        "\\author{" + " \\\\ ".join(
            _tex(str(author).replace("|", ", "))
            for author in state.get("paper_metadata", {}).get("authors", [])
        ) + "}",
        "\\maketitle",
    ]
    if d["abstract"]:
        tex.append(f"\\begin{{abstract}}\n{_tex(d['abstract'])}\n\\end{{abstract}}")
    if d["index_terms"]:
        tex.append(f"\\begin{{IEEEkeywords}}\n{_tex(', '.join(d['index_terms']))}\n\\end{{IEEEkeywords}}")
    for title, text in d["sections"]:
        tex.append(f"\\section{{{_tex(title)}}}\n\n{_tex(text)}")
    tex.append("\\begin{thebibliography}{99}")
    for i, r in enumerate(d["references"], 1):
        tex.append(f"\\bibitem{{ref{i}}} {_tex(r)}")
    tex.append("\\end{thebibliography}")
    tex.append("\\end{document}")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(tex) + "\n")
    return path


def document_export_node(state: ResearchState) -> dict:
    print("\n--- [Stage 17] Document Export ---")
    status = state.get("final_validation_status", "")
    forced = state.get("revision_count", 0) >= MAX_REVISIONS
    if status != "pass" and not forced:
        print("Export skipped: final validation did not pass.")
        return {"output_files": state.get("output_files", [])}
    if status != "pass":
        print("WARNING: exporting a manuscript with unresolved validation issues:")
        for issue in state.get("final_validation_issues", []):
            print(f"  - {issue}")

    output_dir = os.path.join("outputs", "final")
    os.makedirs(output_dir, exist_ok=True)
    files = list(state.get("output_files", []))
    for label, fn in (("DOCX", export_docx), ("PDF", export_pdf), ("LaTeX", export_latex)):
        try:
            p = fn(state, output_dir)
            files.append(p)
            print(f"{label} generated: {p}")
        except Exception as exc:  # noqa: BLE001
            print(f"{label} generation failed: {exc}")
    return {"output_files": files}



# GRAPH


graph = StateGraph(ResearchState)

# Nodes that call an LLM get the retry policy (handles transient 429/5xx).
graph.add_node("supervisor", supervisor_node)
graph.add_node("project_understanding", project_understanding_node, retry_policy=RETRY)
graph.add_node("research_planner", research_planner_node, retry_policy=RETRY)
graph.add_node("literature_research", literature_research_node, retry_policy=RETRY)
graph.add_node("semantic_scholar", ToolNode(tools))
graph.add_node("collect_literature", collect_literature_node)
graph.add_node("paper_filtering", paper_filtering_node)
graph.add_node("literature_analysis", literature_analysis_node, retry_policy=RETRY)
graph.add_node("literature_synthesis", literature_synthesis_node, retry_policy=RETRY)
graph.add_node("research_gap", research_gap_node, retry_policy=RETRY)
graph.add_node("paper_outline", paper_outline_node, retry_policy=RETRY)
graph.add_node("human_approval", human_approval_node)
graph.add_node("writer", writer_node, retry_policy=RETRY)
graph.add_node("citation_verification", citation_verification_node, retry_policy=RETRY)
graph.add_node("claim_verification", claim_verification_node, retry_policy=RETRY)
graph.add_node("consistency_verification", consistency_verification_node, retry_policy=RETRY)
graph.add_node("review", review_node, retry_policy=RETRY)
graph.add_node("revision", revision_node, retry_policy=RETRY)
graph.add_node("formatting", formatting_node, retry_policy=RETRY)
graph.add_node("final_validation", final_validation_node, retry_policy=RETRY)
graph.add_node("document_export", document_export_node)

graph.add_edge(START, "supervisor")
graph.add_edge("supervisor", "project_understanding")
graph.add_edge("project_understanding", "research_planner")
graph.add_edge("research_planner", "literature_research")

graph.add_conditional_edges(
    "literature_research",
    route_after_literature,
    {
        "semantic_scholar": "semantic_scholar",
        "literature_research": "literature_research",
        "collect_literature": "collect_literature",
    },
)
graph.add_edge("semantic_scholar", "literature_research")
graph.add_edge("collect_literature", "paper_filtering")
graph.add_edge("paper_filtering", "literature_analysis")
graph.add_edge("literature_analysis", "literature_synthesis")
graph.add_edge("literature_synthesis", "research_gap")

graph.add_edge("research_gap", "paper_outline")
graph.add_edge("paper_outline", "human_approval")
graph.add_conditional_edges(
    "human_approval",
    route_after_human_approval,
    {"writer": "writer", "research_planner": "research_planner"},
)


graph.add_conditional_edges(
    "writer",
    route_after_writer,
    {"writer": "writer", "citation_verification": "citation_verification"},
)

graph.add_edge("citation_verification", "claim_verification")
graph.add_edge("claim_verification", "consistency_verification")
graph.add_edge("consistency_verification", "review")

graph.add_conditional_edges(
    "review", route_after_review, {"revision": "revision", "formatting": "formatting"}
)
# The approved workflow continues from Revision to Formatting.
graph.add_edge("revision", "formatting")

graph.add_edge("formatting", "final_validation")
graph.add_conditional_edges(
    "final_validation",
    route_after_final_validation,
    {"document_export": "document_export", "revision": "revision"},
)
graph.add_edge("document_export", END)

app = graph.compile(checkpointer=InMemorySaver())



# RUN



def _ask_approval(payload: dict) -> dict:
    print("\n" + "=" * 70)
    print("OUTLINE APPROVAL REQUIRED")
    print("=" * 70)
    titles = payload.get("title_candidates", [])
    for i, t in enumerate(titles, 1):
        print(f"  Title {i}: {t}")
    for s in payload.get("paper_outline", {}).get("sections", []):
        print(f"\n  {s.get('section_number', '')} {s.get('title', '')}: {s.get('purpose', '')}")
        for kp in s.get("key_points", []):
            print(f"     - {kp}")
    print()
    answer = input("Approve outline? [y/n]: ").strip().lower()
    if answer.startswith("y"):
        resume: dict[str, Any] = {"status": "approved", "feedback": ""}
        if titles:
            pick = input(f"Choose title 1-{len(titles)} (Enter = 1): ").strip()
            if pick.isdigit() and 1 <= int(pick) <= len(titles):
                resume["selected_title"] = titles[int(pick) - 1]
        return resume
    return {"status": "rejected", "feedback": input("Feedback for the planner: ").strip()}


if __name__ == "__main__":
    initial_state = {
        "project_description": "Build an AI system that generates research papers from project data and literature.",
        "uploaded_files": [],
        "messages": [],
        
        "actual_results": {},
    }
    config = {"configurable": {"thread_id": "paperforge-001"}, "recursion_limit": 150}

    result = app.invoke(initial_state, config=config)

    # Human-in-the-loop: resume every time the graph pauses at an interrupt.
    while "__interrupt__" in result:
        payload = result["__interrupt__"][0].value
        result = app.invoke(Command(resume=_ask_approval(payload)), config=config)

    print("\nDone.")
    print("Validation status:", result.get("final_validation_status"))
    print("Output files:", result.get("output_files"))